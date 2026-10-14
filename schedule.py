from datetime import datetime, time
from typing import Optional, Dict, List, Tuple
import openpyxl
import logging

logger = logging.getLogger(name)

class ScheduleManager:
    def init(self, file_path: str):
        self.file_path = file_path
        # Облегченный кэш расписания
        self._schedule_cache: Dict[Tuple[str, str], List[Tuple[time, time, str, str]]] = {}
        # Облегченный кэш учителей
        self._teachers_cache: Dict[str, Dict[str, str]] = {}
        self.days_mapping = {
            0: "Понедельник", 1: "Вторник", 2: "Среда",
            3: "Четверг", 4: "Пятница", 5: "Суббота", 6: "Воскресенье"
        }

    def load_schedule(self) -> None:
        try:
            # Открываем файл строго в режиме энергосбережения оперативной памяти
            wb = openpyxl.load_workbook(self.file_path, read_only=True, data_only=True)
            sheet = wb.active
            
            self._schedule_cache.clear()
            self._teachers_cache.clear()
            
            # Разбираем заголовки колонок
            header_row = next(sheet.iter_rows(max_row=1, values_only=True))
            header = [str(cell).strip() if cell else "" for cell in header_row]
            
            try:
                col_class = header.index("Класс")
                col_day = header.index("День недели")
                col_time = header.index("Время")
                col_subj = header.index("Предмет")
                col_teach = header.index("Учитель") if "Учитель" in header else None
            except ValueError as e:
                logger.error(f"Ошибка: Неверная шапка таблицы Excel: {e}")
                wb.close()
                return

            # Читаем все строчки уроков подряд
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if not row or row[col_class] is None:
                    continue
                
                school_class = str(row[col_class]).strip().upper()
                day_of_week = str(row[col_day]).strip().capitalize()
                time_str = str(row[col_time]).strip()
                subject = str(row[col_subj]).strip()
                
                if col_teach is not None and row[col_teach] is not None:
                    teacher = str(row[col_teach]).strip()
                else:
                    teacher = "Не указан"
                
                try:
                    start_str, end_str = time_str.split('-')
                    t_start = datetime.strptime(start_str.strip(), "%H:%M").time()
                    t_end = datetime.strptime(end_str.strip(), "%H:%M").time()
                except Exception as e:
                    logger.error(f"Пропущена строка из-за неверного формата времени '{time_str}': {e}")
                    continue
                
                # Сохраняем урок в расписание
                key = (school_class, day_of_week)
                if key not in self._schedule_cache:
                    self._schedule_cache[key] = []
                self._schedule_cache[key].append((t_start, t_end, subject, teacher))
                
                # Привязываем учителя к предмету для этого конкретного класса
                if school_class not in self._teachers_cache:
                    self._teachers_cache[school_class] = {}
                self._teachers_cache[school_class][subject.lower()] = teacher
            
            wb.close() # Полностью выгружаем Excel-файл из памяти сервера

            # Сортируем время уроков
            for key in self._schedule_cache:
                self._schedule_cache[key].sort(key=lambda x: x[0])
                
            logger.info("Успешно: База расписания оптимизирована и загружена в память через openpyxl!")
        except Exception as e:
            logger.error(f"Критическая ошибка парсера: {e}")

    def get_all_classes(self) -> List[str]:
        return sorted(list(self._teachers_cache.keys()))
        def get_lessons_for_time(self, school_class: str, target_day: str, target_time: time) -> Tuple[Optional[str], Optional[str]]:
        key = (school_class.upper(), target_day.capitalize())
        lessons = self._schedule_cache.get(key, [])
        if not lessons:
            return None, None

        current_lesson, next_lesson = None, None

        for i, (t_start, t_end, subject, teacher) in enumerate(lessons):
            if t_start <= target_time <= t_end:
                current_lesson = f"{subject} ({t_start.strftime('%H:%M')}-{t_end.strftime('%H:%M')}) — {teacher}"
                if i + 1 < len(lessons):
                    n_start, n_end, n_subj, _ = lessons[i+1]
                    next_lesson = f"{n_subj} ({n_start.strftime('%H:%M')}-{n_end.strftime('%H:%M')})"
                break
            elif target_time < t_start:
                next_lesson = f"{subject} ({t_start.strftime('%H:%M')}-{t_end.strftime('%H:%M')})"
                break
                
        return current_lesson, next_lesson

    def get_full_day_schedule(self, school_class: str, target_day: str) -> str:
        key = (school_class.upper(), target_day.capitalize())
        lessons = self._schedule_cache.get(key, [])
        if not lessons:
            return f"На {target_day} расписание для класса {school_class} не найдено."
        
        result = f"📅 Расписание на {target_day} ({school_class}):\n"
        for i, (t_start, t_end, subject, teacher) in enumerate(lessons, 1):
            result += f"{i}. {t_start.strftime('%H:%M')}-{t_end.strftime('%H:%M')} — {subject} ({teacher})\n"
        return result

    def get_subjects_for_class(self, school_class: str) -> List[str]:
        subjects_dict = self._teachers_cache.get(school_class.upper(), {})
        return sorted([sub.capitalize() for sub in subjects_dict.keys()])

    def get_teacher_for_subject(self, school_class: str, subject: str) -> str:
        subjects_dict = self._teachers_cache.get(school_class.upper(), {})
        return subjects_dict.get(subject.lower(), "Учитель не указан")
